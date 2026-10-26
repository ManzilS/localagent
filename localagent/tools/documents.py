"""Documents and vision.

read_document: PDF, DOCX, XLSX/CSV, PPTX, HTML, JSON, text. Scanned PDF pages
(no text layer) are rendered and read by the vision model.
look_at: ask the vision model about an image file or a fresh screenshot.
"""
from __future__ import annotations

import csv
import json
import tempfile
from datetime import datetime
from pathlib import Path

from ..safety import resolve_path
from . import tool, truncate

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}


def _pages(spec: str | None, n: int) -> list[int]:
    if not spec:
        return list(range(min(n, 30)))
    out: list[int] = []
    for part in str(spec).split(","):
        part = part.strip()
        if "-" in part:
            a, b = part.split("-", 1)
            out += list(range(int(a) - 1, min(int(b), n)))
        elif part:
            out.append(int(part) - 1)
    return [p for p in out if 0 <= p < n]


def ask_vision(rt, images: list[str], question: str) -> str:
    prov = rt.vision_provider()
    if prov is None:
        return "ERROR: no vision model configured (models.vision in config.yaml)"
    resp = prov.chat([
        {"role": "system", "content": "You are a precise visual analyst. Describe exactly what is shown. "
                                      "Transcribe all visible text verbatim when asked or when it matters."},
        {"role": "user", "content": question, "images": images},
    ], temperature=0.1, max_tokens=2048, num_ctx=rt.ctx_for(prov))
    return resp.content.strip() or "(vision model returned nothing)"


def _render_pdf_page(path: Path, index: int) -> str | None:
    try:
        import pypdfium2 as pdfium  # type: ignore
    except ImportError:
        return None
    pdf = pdfium.PdfDocument(str(path))
    img = pdf[index].render(scale=1.6).to_pil()
    out = Path(tempfile.gettempdir()) / f"la_pdf_{path.stem}_{index + 1}.png"
    img.save(out)
    return str(out)


def _read_pdf(rt, p: Path, pages: str | None) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(p))
    n = len(reader.pages)
    chunks = [f"[PDF: {n} pages]"]
    scanned = 0
    for i in _pages(pages, n):
        text = (reader.pages[i].extract_text() or "").strip()
        if len(text) < 25:
            img = _render_pdf_page(p, i)
            if img and scanned < 6:
                scanned += 1
                text = "[scanned page, read by vision model]\n" + ask_vision(
                    rt, [img], "Transcribe all text on this page verbatim, preserving tables as markdown.")
            elif not img:
                text = text or "[no text layer — install `pip install localagent[vision]` to OCR scanned pages]"
        chunks.append(f"--- page {i + 1} ---\n{text}")
    if not pages and n > 30:
        chunks.append(f"[stopped at page 30 of {n}; pass pages='31-60' for more]")
    return "\n\n".join(chunks)


def _read_docx(p: Path) -> str:
    import docx  # python-docx
    d = docx.Document(str(p))
    parts = []
    for para in d.paragraphs:
        if para.text.strip():
            style = (para.style.name or "").lower() if para.style is not None else ""
            prefix = "#" * int(style[-1]) + " " if style.startswith("heading") and style[-1:].isdigit() else ""
            parts.append(prefix + para.text)
    for ti, table in enumerate(d.tables, 1):
        parts.append(f"\n[table {ti}]")
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _read_xlsx(p: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"## sheet: {ws.title}")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= 300:
                parts.append("[... more rows; use run_python with openpyxl/pandas for full analysis]")
                break
            if any(v is not None for v in row):
                parts.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(parts)


def _read_pptx(p: Path) -> str:
    try:
        from pptx import Presentation  # type: ignore
    except ImportError:
        return "[install python-pptx (`pip install localagent[office]`) to read .pptx]"
    prs = Presentation(str(p))
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        texts = [sh.text_frame.text for sh in slide.shapes if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip()]
        parts.append(f"--- slide {i} ---\n" + "\n".join(texts))
    return "\n\n".join(parts)


@tool("read_document", "Extract text from a document: PDF (incl. scanned), Word .docx, Excel .xlsx, .csv, "
      "PowerPoint .pptx, HTML, JSON or text. For images use look_at.",
      {"path": {"type": "string"}, "pages": {"type": "string", "description": "PDF pages like '1-5,9'"}}, ["path"], group="vision")
def read_document(rt, path: str, pages: str | None = None) -> str:
    p = resolve_path(rt.cfg, path)
    ext = p.suffix.lower()
    if ext == ".pdf":
        text = _read_pdf(rt, p, pages)
    elif ext == ".docx":
        text = _read_docx(p)
    elif ext in (".xlsx", ".xlsm"):
        text = _read_xlsx(p)
    elif ext == ".csv":
        with open(p, newline="", encoding="utf-8", errors="replace") as f:
            rows = [" | ".join(r) for _, r in zip(range(300), csv.reader(f))]
        text = "\n".join(rows)
    elif ext == ".pptx":
        text = _read_pptx(p)
    elif ext in (".html", ".htm"):
        from .web import html_to_text
        text = html_to_text(p.read_text(encoding="utf-8", errors="replace"))
    elif ext == ".json":
        text = json.dumps(json.loads(p.read_text(encoding="utf-8")), indent=2)
    elif ext in IMAGE_EXT:
        return look_at(rt, str(p), "Describe this image and transcribe any text in it.")
    else:
        text = p.read_text(encoding="utf-8", errors="replace")
    return truncate(text, rt.cfg.agent["max_tool_output_chars"])


def take_screenshot(rt) -> str:
    try:
        import mss  # type: ignore
        import mss.tools  # type: ignore
    except ImportError:
        raise RuntimeError("screenshots need `pip install localagent[vision]` (mss)")
    out_dir = rt.cfg.workspace / "screenshots"
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"screen_{datetime.now():%Y%m%d_%H%M%S}.png"
    with mss.mss() as s:
        shot = s.grab(s.monitors[1])
        mss.tools.to_png(shot.rgb, shot.size, output=str(out))
    return str(out)


@tool("look_at", "Ask the vision model about an image: photos, screenshots, charts, diagrams, handwriting, UI mockups. "
      "Use path='screen' to capture the current screen first.",
      {"path": {"type": "string", "description": "image file, or 'screen'"},
       "question": {"type": "string", "description": "what you want to know"}}, ["path"], group="vision")
def look_at(rt, path: str, question: str = "Describe this image in detail, including any text.") -> str:
    if path.strip().lower() in ("screen", "screenshot"):
        img = take_screenshot(rt)
        return f"[screenshot saved to {img}]\n" + ask_vision(rt, [img], question)
    p = resolve_path(rt.cfg, path)
    if p.suffix.lower() not in IMAGE_EXT:
        return f"{p.name} is not an image; use read_document"
    return ask_vision(rt, [str(p)], question)
